#!/usr/bin/env python3
"""Recompute manuscript Table Panel B for the automatic misstep detector.

Inputs:
- automatic detector output in misstep_corpus_app_results.json
- the two completed blinded clinician workbooks in ../completed/

The completed workbooks have an invalid core-properties XML namespace in the
archived files. They are repaired in memory before openpyxl reads them; the
original xlsx files are not modified.
"""

from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
MISSTEP_DIR = SCRIPT_DIR.parent
COMPLETED_DIR = MISSTEP_DIR / "completed"
RESULTS_PATH = SCRIPT_DIR / "misstep_corpus_app_results.json"
OUTPUT_PATH = SCRIPT_DIR / "panel_b_metrics.csv"

RATER_1_PATH = COMPLETED_DIR / "LLMPatients_Misstep_Clinician_Evaluation_Rater_El_Completed.xlsx"
RATER_2_PATH = COMPLETED_DIR / "LLMPatients_Misstep_Clinician_Evaluation_Rater_Er_Completed.xlsx"

EXPECTED_PANEL_B = {
    "Rater 1": (159, 34, 44, 63, 0.824, 0.716, 0.766, 0.677),
    "Rater 2": (134, 59, 86, 21, 0.694, 0.865, 0.770, 0.733),
    "Either clinician positive": (161, 32, 44, 63, 0.834, 0.719, 0.772, 0.683),
    "Both clinicians positive": (132, 61, 86, 21, 0.684, 0.863, 0.763, 0.727),
    "Therapist turns with deliberately scripted missteps": (
        131,
        62,
        88,
        19,
        0.679,
        0.873,
        0.764,
        0.730,
    ),
}


def repair_core_properties_xml(raw: bytes) -> bytes:
    text = raw.decode("utf-8")
    if "xmlns:dc=" not in text:
        text = text.replace(
            "<cp:coreProperties ",
            '<cp:coreProperties xmlns:dc="http://purl.org/dc/elements/1.1/" ',
            1,
        )
    return text.encode("utf-8")


def load_workbook_repaired(path: Path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(path, "r") as source:
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "docProps/core.xml":
                    data = repair_core_properties_xml(data)
                target.writestr(info, data)
    buffer.seek(0)
    return load_workbook(buffer, data_only=True, read_only=True)


def yes(value: object) -> bool:
    return str(value or "").strip().lower() == "yes"


def read_rater_labels(path: Path) -> dict[tuple[str, int], bool]:
    workbook = load_workbook_repaired(path)
    labels: dict[tuple[str, int], bool] = {}
    for sheet_name in sorted(name for name in workbook.sheetnames if name.startswith("T")):
        worksheet = workbook[sheet_name]
        for row_idx in range(16, 31):
            therapist_turn = row_idx - 15
            labels[(sheet_name, therapist_turn)] = yes(worksheet.cell(row_idx, 5).value)
    return labels


def read_detector_predictions() -> tuple[list[tuple[str, int]], dict[tuple[str, int], bool], dict[tuple[str, int], bool]]:
    results = json.loads(RESULTS_PATH.read_text(encoding="utf-8"))
    units: list[tuple[str, int]] = []
    predicted: dict[tuple[str, int], bool] = {}
    seeded_reference: dict[tuple[str, int], bool] = {}

    for result in results:
        transcript_id = result["transcriptId"]
        therapist_turn_count = int(result["summary"]["therapistTurnCount"])
        flagged_turns: set[int] = set()
        for category in result["detectedCategories"]:
            flagged_turns.update(int(turn) for turn in category.get("therapistTurns", []))

        is_seeded = result["scriptVariant"] == "seeded_misstep"
        for turn in range(1, therapist_turn_count + 1):
            unit = (transcript_id, turn)
            units.append(unit)
            predicted[unit] = turn in flagged_turns
            seeded_reference[unit] = is_seeded

    return units, predicted, seeded_reference


def metrics(
    units: list[tuple[str, int]],
    predicted: dict[tuple[str, int], bool],
    reference: Callable[[tuple[str, int]], bool],
) -> tuple[int, int, int, int, float, float, float, float]:
    tp = fp = tn = fn = 0
    for unit in units:
        pred = predicted[unit]
        ref = reference(unit)
        if pred and ref:
            tp += 1
        elif pred and not ref:
            fp += 1
        elif not pred and not ref:
            tn += 1
        else:
            fn += 1

    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    accuracy = (tp + tn) / len(units) if units else 0.0
    return tp, fp, tn, fn, precision, recall, f1, accuracy


def rounded(row: tuple[int, int, int, int, float, float, float, float]):
    tp, fp, tn, fn, precision, recall, f1, accuracy = row
    return (tp, fp, tn, fn, round(precision, 3), round(recall, 3), round(f1, 3), round(accuracy, 3))


def csv_row(row: tuple[str, int, int, int, int, float, float, float, float]):
    label, tp, fp, tn, fn, precision, recall, f1, accuracy = row
    return [label, tp, fp, tn, fn, f"{precision:.3f}", f"{recall:.3f}", f"{f1:.3f}", f"{accuracy:.3f}"]


def main() -> None:
    units, predicted, seeded_reference = read_detector_predictions()
    rater_1 = read_rater_labels(RATER_1_PATH)
    rater_2 = read_rater_labels(RATER_2_PATH)

    comparisons = {
        "Rater 1": lambda unit: rater_1[unit],
        "Rater 2": lambda unit: rater_2[unit],
        "Either clinician positive": lambda unit: rater_1[unit] or rater_2[unit],
        "Both clinicians positive": lambda unit: rater_1[unit] and rater_2[unit],
        "Therapist turns with deliberately scripted missteps": lambda unit: seeded_reference[unit],
    }

    rows = []
    for label, reference in comparisons.items():
        row = rounded(metrics(units, predicted, reference))
        if row != EXPECTED_PANEL_B[label]:
            raise SystemExit(f"{label} mismatch: computed {row}, expected {EXPECTED_PANEL_B[label]}")
        rows.append((label, *row))

    with OUTPUT_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["reference", "tp", "fp", "tn", "fn", "precision", "recall", "f1", "accuracy"])
        writer.writerows(csv_row(row) for row in rows)

    print(f"Wrote {OUTPUT_PATH}")
    for row in rows:
        label, tp, fp, tn, fn, precision, recall, f1, accuracy = row
        print(f"{label}: TP={tp} FP={fp} TN={tn} FN={fn} P={precision:.3f} R={recall:.3f} F1={f1:.3f} Acc={accuracy:.3f}")


if __name__ == "__main__":
    main()
